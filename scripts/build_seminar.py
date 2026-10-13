#!/usr/bin/env python3
"""Build a fully editable Persian seminar with verifiable references and charts.
Run from any directory: python scripts/build_seminar.py [--font-size 12.5]
No downloaded fonts or copyrighted article texts are bundled.
"""
from __future__ import annotations
import argparse, csv, json, re, sys
from pathlib import Path
from collections import Counter
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from content_assets import REFERENCES, TABLES, FIGURES, ABSTRACT

ROOT=Path(__file__).resolve().parents[1]
OUT=ROOT/'artifacts'; OUT.mkdir(exist_ok=True)
FAFONT='Noto Naskh Arabic'
ENFONT='Times New Roman'
FN_TRANS=str.maketrans('0123456789','۰۱۲۳۴۵۶۷۸۹')
def fa(x): return str(x).translate(FN_TRANS)
CITE_RE=re.compile(r'\[@([^\]]+)\]')
TABLE_MAP={1:1,2:2,3:3,6:4,4:5,5:6}
used=set(); first_citation=set(); caption_info=[]; headings=[]


def xml(tag,**attrs):
    e=OxmlElement(tag)
    for k,v in attrs.items(): e.set(qn(k),str(v))
    return e

def direction(p,rtl=True):
    pp=p._p.get_or_add_pPr(); old=pp.find(qn('w:bidi'))
    if old is None: pp.append(xml('w:bidi',**{'w:val':int(rtl)}))
    else: old.set(qn('w:val'),str(int(rtl)))
    return p

def font(run,size=None,rtl=True,bold=None,italic=None):
    rp=run._r.get_or_add_rPr()
    rf=rp.find(qn('w:rFonts'))
    if rf is None:rf=xml('w:rFonts');rp.insert(0,rf)
    for key,val in [('ascii',ENFONT),('hAnsi',ENFONT),('cs',FAFONT),('eastAsia',FAFONT)]:rf.set(qn('w:'+key),val)
    r=rp.find(qn('w:rtl'))
    if r is None:rp.append(xml('w:rtl',**{'w:val':int(rtl)}))
    else:r.set(qn('w:val'),str(int(rtl)))
    lang=rp.find(qn('w:lang'))
    if lang is None:rp.append(xml('w:lang',**{'w:val':'en-US','w:bidi':'fa-IR'}))
    if size is not None:
        run.font.size=Pt(size)
        el=rp.find(qn('w:szCs'))
        if el is None:rp.append(xml('w:szCs',**{'w:val':int(size*2)}))
        else:el.set(qn('w:val'),str(int(size*2)))
    if bold is not None:
        run.bold=bold; e=rp.find(qn('w:bCs'))
        if e is None:rp.append(xml('w:bCs',**{'w:val':int(bold)}))
        else:e.set(qn('w:val'),str(int(bold)))
    if italic is not None:run.italic=italic
    return run

def cite_replace(m):
    ids=m.group(1).split(';')
    ids=sorted(ids,key=lambda k:REFERENCES[k]['sort'].lower())
    vals=[]
    for k in ids:
        if k not in REFERENCES: raise ValueError('Unknown citation '+k)
        rr=REFERENCES[k];used.add(k)
        v=rr.get('first_cite',rr['cite']) if k not in first_citation else rr['cite']
        first_citation.add(k);vals.append(v.replace(' ', '\u00a0'))
    return '('+'; '.join(vals)+')'

LATIN=re.compile(r'(\([A-Za-z][^()\n\u0600-\u06ff]*\)|[A-Za-z0-9][A-Za-z0-9 \[\]().,;:/%+\-–—&=’\'‐−]*[A-Za-z0-9\])]|[A-Za-z])')
def rich(p,text,size=None,bold=None):
    s=CITE_RE.sub(cite_replace,text)
    cur=0
    for m in LATIN.finditer(s):
        if m.start()>cur:font(p.add_run(s[cur:m.start()]),size,True,bold)
        font(p.add_run('\u200e'+m.group(0)+'\u200e'),size,False,bold)
        cur=m.end()
    if cur<len(s):font(p.add_run(s[cur:]),size,True,bold)
    return p

def para(doc,text='',style=None,size=None,align=None):
    p=doc.add_paragraph(style=style)
    direction(p)
    if align is not None:p.alignment=align
    rich(p,text,size)
    return p

def field(p,instruction,cached=''):
    r=p.add_run();font(r)
    r._r.append(xml('w:fldChar',**{'w:fldCharType':'begin'}))
    t=xml('w:instrText',**{'xml:space':'preserve'});t.text=' '+instruction+' ';r._r.append(t)
    r._r.append(xml('w:fldChar',**{'w:fldCharType':'separate'}))
    if cached:font(p.add_run(cached),rtl=True)
    r2=p.add_run();r2._r.append(xml('w:fldChar',**{'w:fldCharType':'end'}))

def bookmark(p,name):
    i=len(headings)+len(caption_info)+100
    p._p.insert(0,xml('w:bookmarkStart',**{'w:id':i,'w:name':name}))
    p._p.append(xml('w:bookmarkEnd',**{'w:id':i}))

def add_heading(doc,text,level):
    p=para(doc,text,'Heading '+str(level))
    for r in p.runs:font(r,bold=True)
    name=f'h{len(headings)+1}'
    bookmark(p,name);headings.append(dict(text=text,level=level,bookmark=name))
    return p

def caption(doc,kind,num,title):
    style='TableCaption' if kind=='Table' else 'FigureCaption'
    p=doc.add_paragraph(style=style);direction(p);p.paragraph_format.keep_with_next=True
    rich(p,'جدول ' if kind=='Table' else 'شکل ',11.5,bold=True)
    field(p,f'SEQ {kind} \\* ARABIC',fa(num))
    rich(p,'. '+title,11.5,bold=True)
    bookmark(p,f'{kind}{num}')
    caption_info.append(dict(kind=kind,number=num,title=title))
    return p

def cell_borders(table):
    pr=table._tbl.tblPr
    borders=xml('w:tblBorders')
    for pos in ['top','left','bottom','right','insideH','insideV']:
        borders.append(xml('w:'+pos,**{'w:val':'single','w:sz':'4','w:color':'B7B7B7'}))
    pr.append(borders)
    margins=xml('w:tblCellMar')
    for pos in ['top','bottom']:margins.append(xml('w:'+pos,**{'w:w':'70','w:type':'dxa'}))
    for pos in ['left','right']:margins.append(xml('w:'+pos,**{'w:w':'90','w:type':'dxa'}))
    pr.append(margins)
    pr.append(xml('w:bidiVisual',**{'w:val':'1'}))

def table_block(doc,n):
    d=TABLES[n];num=TABLE_MAP[n]
    caption(doc,'Table',num,d['title'])
    t=doc.add_table(rows=1,cols=len(d['heads']))
    t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False;cell_borders(t)
    for col,w in zip(t.columns,d['widths']):col.width=Cm(w)
    for c,h,w in zip(t.rows[0].cells,d['heads'],d['widths']):
        c.width=Cm(w);c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
        c._tc.get_or_add_tcPr().append(xml('w:shd',**{'w:fill':'EAEAEA'}))
        p=c.paragraphs[0];p.style='TableText';direction(p);rich(p,h,10.5,True)
    t.rows[0]._tr.get_or_add_trPr().append(xml('w:tblHeader',**{'w:val':'1'}))
    for rr in d['rows']:
        row=t.add_row();row._tr.get_or_add_trPr().append(xml('w:cantSplit'))
        for c,txt,w in zip(row.cells,rr,d['widths']):
            c.width=Cm(w);c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p=c.paragraphs[0];p.style='TableText';direction(p);rich(p,txt,10.5)
    para(doc,'یادداشت. '+d['note'],'NoteText')


def plot_charts():
    with (ROOT/'research/data/effect_estimates.csv').open() as f:rows=list(csv.DictReader(f))
    configs=[('Batty2018',2,'Adjusted hazard ratio per 1 SD',1,(0.45,1.7)),('Yen2021',3,'Adjusted odds ratio',1,(0.65,3.95)),('Moore2022',4,'Pooled correlation (r)',0,(-.025,.345))]
    for sid,num,xlabel,null,xlim in configs:
        r=[x for x in rows if x['study']==sid]
        est=[float(x['estimate']) for x in r]; lo=[float(x['lower']) for x in r];hi=[float(x['upper']) for x in r]
        y=list(range(len(r)))
        fig,ax=plt.subplots(figsize=(7.8,3.15 if len(r)==3 else 3.8))
        ax.errorbar(est,y,xerr=[[v-l for v,l in zip(est,lo)],[h-v for h,v in zip(hi,est)]],fmt='o',capsize=4,markersize=5)
        ax.axvline(null,linestyle='--',linewidth=1)
        ax.set_yticks(y,[x['construct'] for x in r]);ax.invert_yaxis();ax.set_xlim(*xlim);ax.set_xlabel(xlabel,fontsize=10)
        ax.spines['top'].set_visible(False);ax.spines['right'].set_visible(False)
        ax.tick_params(axis='both',labelsize=9)
        for i,(e,l,h) in enumerate(zip(est,lo,hi)):
            ax.text(xlim[1],i-.20,f'{e:.2f} [{l:.2f}, {h:.2f}]',ha='right',va='bottom',fontsize=8)
        fig.tight_layout(pad=1.2)
        fig.savefig(OUT/f'figure_{num}.png',dpi=210)
        plt.close(fig)
    return rows


def concept(doc):
    # Native Word cells keep this conceptual illustration editable.
    t=doc.add_table(rows=1,cols=2);t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
    for c in t.columns:c.width=Cm(8)
    cell_borders(t)
    levels=[['صفات و الگوهای شخصیت','فشارها، زمینه فرهنگی و شرایط زندگی'],['↓ مسیرهای پیشنهادی ↓'],['ارزیابی شکست و آینده؛ تنظیم هیجان؛ تجربه روابط'],['↓ فرضیه‌های نیازمند آزمون ↓'],['افکار خودکشی'],['↓ گذار احتمالی، نه اجتناب‌ناپذیر ↓'],['اقدام و مرگ: دو پیامد متمایز؛ نیازمند بررسی جداگانه'],['منابع حمایت و مراقبت باید در همه لایه‌ها بررسی شوند.']]
    for idx,row_text in enumerate(levels):
        row=t.rows[0] if idx==0 else t.add_row();row._tr.get_or_add_trPr().append(xml('w:cantSplit'))
        cells=row.cells
        if len(row_text)==1:cells=[cells[0].merge(cells[1])]
        for c,txt in zip(cells,row_text):
            p=c.paragraphs[0];p.style='TableText';p.alignment=WD_ALIGN_PARAGRAPH.CENTER;direction(p);rich(p,txt,11.5,idx in [0,2,4,6])
            if idx in [0,2,4,6]:c._tc.get_or_add_tcPr().append(xml('w:shd',**{'w:fill':'F0F0F0'}))


def figure_block(doc,n):
    d=FIGURES[n];caption(doc,'Figure',n,d['title'])
    if n==1:concept(doc)
    else:
        p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.space_after=Pt(2)
        r=p.add_run();r.add_picture(str(OUT/f'figure_{n}.png'),width=Cm(16))
        pr=r._r.xpath('.//wp:docPr')
        if pr:pr[0].set('descr',d['title'])
    para(doc,'یادداشت. '+d['note'],'NoteText')


def hyperlink(p,label,url,rtl=False,size=10.5):
    rid=p.part.relate_to(url,RT.HYPERLINK,is_external=True)
    h=xml('w:hyperlink',**{'r:id':rid})
    r=p.add_run(label);font(r,size,rtl)
    r.font.underline=True
    p._p.remove(r._r);h.append(r._r);p._p.append(h)


def setup(doc,base):
    sec=doc.sections[0];sec.page_width=Cm(21);sec.page_height=Cm(29.7)
    sec.top_margin=Cm(2.5);sec.bottom_margin=Cm(2.5);sec.left_margin=Cm(2.5);sec.right_margin=Cm(2.5)
    sec.header_distance=Cm(1.2);sec.footer_distance=Cm(1.2);sec.different_first_page_header_footer=True
    normal=doc.styles['Normal'];normal.font.name=ENFONT;normal.font.size=Pt(base)
    normal.paragraph_format.line_spacing=1.5
    normal.paragraph_format.space_after=Pt(1)
    normal.paragraph_format.first_line_indent=Cm(.45)
    normal.paragraph_format.widow_control=True
    normal.paragraph_format.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
    def sfont(s,size,bold=False):
        r=s.element.get_or_add_rPr();rf=r.find(qn('w:rFonts'))
        if rf is None:rf=xml('w:rFonts');r.insert(0,rf)
        for key,val in [('ascii',ENFONT),('hAnsi',ENFONT),('cs',FAFONT)]:rf.set(qn('w:'+key),val)
        for stale in list(r):
            if stale.tag in [qn('w:szCs'),qn('w:bCs')]:r.remove(stale)
        for theme in ['asciiTheme','hAnsiTheme','cstheme','eastAsiaTheme']:rf.attrib.pop(qn('w:'+theme),None)
        r.append(xml('w:szCs',**{'w:val':int(size*2)}));r.append(xml('w:lang',**{'w:val':'en-US','w:bidi':'fa-IR'}))
        r.append(xml('w:bCs',**{'w:val':int(bold)}));s.font.size=Pt(size);s.font.color.rgb=RGBColor(0,0,0)
    sfont(normal,base)
    for level,sz in [(1,15),(2,13.5),(3,12.5)]:
        st=doc.styles['Heading '+str(level)];sfont(st,sz,True);st.font.bold=True
        st.paragraph_format.first_line_indent=Cm(0);st.paragraph_format.alignment=WD_ALIGN_PARAGRAPH.RIGHT
        st.paragraph_format.space_before=Pt(8 if level==1 else 5);st.paragraph_format.space_after=Pt(3)
        st.paragraph_format.keep_with_next=True;st.paragraph_format.line_spacing=1.2
    for name,size,spacing in [('TableCaption',11.5,1.1),('FigureCaption',11.5,1.1),('TableText',10.5,1.15),('NoteText',10.5,1.15),('FrontTitle',14,1.2),('TOC 1',11,1.1),('TOC 2',10.5,1.1)]:
        st=doc.styles[name] if name in doc.styles else doc.styles.add_style(name,WD_STYLE_TYPE.PARAGRAPH)
        st.base_style=normal;sfont(st,size,name in ['TableCaption','FigureCaption','FrontTitle'])
        st.paragraph_format.first_line_indent=Cm(0);st.paragraph_format.line_spacing=spacing
        st.paragraph_format.space_before=Pt(4 if 'Caption' in name else 0);st.paragraph_format.space_after=Pt(3)
        st.paragraph_format.alignment=WD_ALIGN_PARAGRAPH.RIGHT
    for name in ['TOC 1','TOC 2']:
        doc.styles[name].paragraph_format.space_after=Pt(2)
    for name in ['TableCaption','FigureCaption']:
        doc.styles[name].paragraph_format.keep_with_next=True
    h=sec.header.paragraphs[0];direction(h);h.alignment=WD_ALIGN_PARAGRAPH.RIGHT
    rich(h,'شخصیت و خودکشی  |  سمینار کارشناسی ارشد روان‌شناسی شخصیت',9)
    ft=sec.footer.paragraphs[0];direction(ft);ft.alignment=WD_ALIGN_PARAGRAPH.CENTER;field(ft,'PAGE','۱')
    settings=doc.settings.element
    settings.append(xml('w:updateFields',**{'w:val':'true'}))
    doc.core_properties.title='رابطه شخصیت و خودکشی: تحلیل نظری و نقد شواهد'
    doc.core_properties.subject='سمینار کارشناسی ارشد روان‌شناسی شخصیت'
    doc.core_properties.author='';doc.core_properties.last_modified_by=''
    doc.core_properties.keywords='شخصیت، خودکشی، مرور انتقادی، پنج‌عاملی، شخصیت مرزی'
    doc.core_properties.comments=''


def add_reference(doc,key):
    r=REFERENCES[key];p=doc.add_paragraph();direction(p,r['lang']=='fa')
    p.paragraph_format.first_line_indent=Cm(-.65)
    p.paragraph_format.left_indent=Cm(.65) if r['lang']=='en' else Cm(0)
    p.paragraph_format.right_indent=Cm(.65) if r['lang']=='fa' else Cm(0)
    p.paragraph_format.line_spacing=1.15;p.paragraph_format.space_after=Pt(8)
    p.paragraph_format.keep_together=True
    p.alignment=WD_ALIGN_PARAGRAPH.LEFT if r['lang']=='en' else WD_ALIGN_PARAGRAPH.RIGHT
    rt=r['lang']=='fa';sz=11 if rt else 10.5
    font(p.add_run(f"{r['authors']} ({r['year']}). {r['title']} "),sz,rt)
    if r['journal']:
        font(p.add_run(r['journal']+', '+r['volume']),sz,rt,italic=True)
        if r['issue']:font(p.add_run('('+r['issue']+')'),sz,rt)
        font(p.add_run(', '+r['pages']+'. '),sz,rt)
    url='https://doi.org/'+r['doi'] if r['doi'] else r['source']
    hyperlink(p,url,url,False,10.5)


def sec_header(doc):
    return doc.sections[0].header._element

def main():
    ap=argparse.ArgumentParser();ap.add_argument('--font-size',type=float,default=11.5);args=ap.parse_args()
    rows=plot_charts();doc=Document();setup(doc,args.font_size)
    # Independent title page.
    p=para(doc,'سمینار کارشناسی ارشد','FrontTitle',align=WD_ALIGN_PARAGRAPH.CENTER)
    p.paragraph_format.space_before=Pt(65)
    para(doc,'روان‌شناسی شخصیت','FrontTitle',align=WD_ALIGN_PARAGRAPH.CENTER)
    p=para(doc,'بررسی جامع رابطه ویژگی‌های شخصیت\nبا افکار، اقدام و رفتارهای خودکشی‌گرایانه',align=WD_ALIGN_PARAGRAPH.CENTER,size=21)
    p.paragraph_format.first_line_indent=Cm(0);p.paragraph_format.space_before=Pt(25);p.paragraph_format.line_spacing=1.2
    para(doc,'تحلیل نظری، شواهد تجربی و بررسی انتقادی سه پژوهش علمی',size=14,align=WD_ALIGN_PARAGRAPH.CENTER)
    p=para(doc,'تهیه‌کننده: ........................................\nاستاد درس: ........................................\nدانشگاه: ........................................',align=WD_ALIGN_PARAGRAPH.CENTER,size=13)
    p.paragraph_format.space_before=Pt(45);p.paragraph_format.first_line_indent=Cm(0)
    para(doc,'۲۰۲۶',align=WD_ALIGN_PARAGRAPH.CENTER,size=12)
    doc.add_page_break()
    para(doc,'۱. صفحات مقدماتی','FrontTitle')
    para(doc,'چکیده','FrontTitle')
    para(doc,ABSTRACT)
    para(doc,'کلیدواژه‌ها: شخصیت؛ افکار خودکشی؛ اقدام به خودکشی؛ مدل پنج‌عاملی؛ شخصیت مرزی؛ تکانشگری؛ نقد روش‌شناختی.','NoteText')
    p=para(doc,'یادداشت خواننده: این متن مرور انتقادی دانشگاهی است؛ مدل مفهومی آن فرضیه‌ساز است و نمره یا الگوریتمی برای تعیین آینده فرد ارائه نمی‌کند.','NoteText')
    p.paragraph_format.space_before=Pt(12)
    doc.add_page_break()
    para(doc,'فهرست مطالب','FrontTitle')
    p=doc.add_paragraph();direction(p);field(p,'TOC \\o "1-2" \\h \\z \\u','فهرست مطالب')
    para(doc,'فهرست جدول‌ها','FrontTitle');p=doc.add_paragraph();direction(p);field(p,'TOC \\h \\z \\t "TableCaption,1"','فهرست جدول‌ها')
    para(doc,'فهرست شکل‌ها و نمودارها','FrontTitle');p=doc.add_paragraph();direction(p);field(p,'TOC \\h \\z \\t "FigureCaption,1"','فهرست شکل‌ها')
    doc.add_page_break()
    src=(ROOT/'manuscript/seminar.fa.md').read_text().replace('جدول‌های ۴ و ۵','جدول‌های ۵ و ۶')
    for block in re.split(r'\n\s*\n',src.strip()):
        block=block.strip()
        if not block:continue
        if block.startswith('#'):
            mark,txt=block.split(' ',1);add_heading(doc,txt,len(mark))
        elif re.fullmatch(r'\[\[TABLE:\d+\]\]',block):table_block(doc,int(re.search(r'\d+',block).group()))
        elif re.fullmatch(r'\[\[FIGURE:\d+\]\]',block):figure_block(doc,int(re.search(r'\d+',block).group()))
        else:para(doc,block)
    p=add_heading(doc,'۱۱. منابع',1);p.paragraph_format.page_break_before=True
    para(doc,'۱۱ـ۱. منابع فارسی','FrontTitle')
    for k in sorted(used,key=lambda k:REFERENCES[k]['sort'].lower()):
        if REFERENCES[k]['lang']=='fa':add_reference(doc,k)
    para(doc,'۱۱ـ۲. منابع انگلیسی','FrontTitle')
    for k in sorted(used,key=lambda k:REFERENCES[k]['sort'].lower()):
        if REFERENCES[k]['lang']=='en':add_reference(doc,k)
    # Logical start prevents LibreOffice from mirroring physical right alignment in RTL.
    for part in [doc.element,doc.styles.element,sec_header(doc)]:
        for jc in part.iter(qn('w:jc')):
            if jc.get(qn('w:val'))=='right':jc.set(qn('w:val'),'start')
    path=OUT/'Personality_and_Suicide_Seminar.fa.docx';doc.save(path)
    unlisted=set(REFERENCES)-used
    report={'font':FAFONT,'body_font_size_pt':args.font_size,'line_spacing':1.5,'manuscript_words':len(src.split()),'references_count':len(used),'uncited_references':sorted(unlisted),'numbered_tables':6,'figures':4,'charts':3,'source_estimates':len(rows),'headings':headings,'captions':caption_info,'docx':path.name}
    (ROOT/'qa/build_report.json').write_text(json.dumps(report,ensure_ascii=False,indent=2))
    if unlisted:raise RuntimeError('Uncited references: '+str(unlisted))
    print(json.dumps({k:v for k,v in report.items() if k not in ['headings','captions']},ensure_ascii=False,indent=2))

if __name__=='__main__':main()
